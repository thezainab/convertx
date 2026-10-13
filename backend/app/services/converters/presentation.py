import os
import shutil
import subprocess
import tempfile
from pptx import Presentation
from app.services.conversion_registry import registry


def _find_libreoffice():
    return shutil.which("libreoffice") or shutil.which("soffice")


def _lo_convert(input_path: str, output_path: str, fmt: str) -> bool:
    """LibreOffice se file convert karta hai aur output ko expected naam par rakhta hai."""
    try:
        in_path = os.path.abspath(input_path)
        out_path = os.path.abspath(output_path)

        lo = _find_libreoffice()
        if not lo:
            print("LibreOffice is server par installed nahi hai")
            return False

        with tempfile.TemporaryDirectory() as tmp_dir:
            result = subprocess.run(
                [lo, "--headless", "--norestore", "--convert-to", fmt,
                 "--outdir", tmp_dir, in_path],
                capture_output=True,
                text=True,
                timeout=120,
                env={**os.environ, "HOME": "/tmp"},
            )
            if result.returncode != 0:
                print(f"LibreOffice error: {result.stderr}")
                return False

            produced = os.path.join(
                tmp_dir, os.path.splitext(os.path.basename(in_path))[0] + "." + fmt
            )
            if not os.path.exists(produced):
                print(f"LibreOffice ne output nahi banaya: {result.stdout}")
                return False

            shutil.move(produced, out_path)
        return True
    except subprocess.TimeoutExpired:
        print("LibreOffice conversion timeout ho gayi")
        return False
    except Exception as e:
        print(f"LibreOffice conversion error: {e}")
        return False


def pptx_to_txt(input_path: str, output_path: str) -> bool:
    try:
        prs = Presentation(input_path)
        with open(output_path, "w", encoding="utf-8") as f:
            for i, slide in enumerate(prs.slides):
                f.write(f"--- Slide {i+1} ---\n")
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                        f.write(shape.text_frame.text + "\n")
        return True
    except Exception as e:
        print(f"PPTX to TXT error: {e}")
        return False


def pptx_to_pdf(input_path: str, output_path: str) -> bool:
    # Asli PDF LibreOffice se banti hai (pptx aur ppt dono ke liye)
    return _lo_convert(input_path, output_path, "pdf")


registry.register('pptx', 'txt', pptx_to_txt)
registry.register('pptx', 'pdf', pptx_to_pdf)
registry.register('ppt', 'pdf', pptx_to_pdf)